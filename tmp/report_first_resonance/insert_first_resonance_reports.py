from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu


YIFENG_DOC = Path(r"E:\量子散射\Yifeng CHEN 999016959\Quantum Scattering Project Report.docx")
YIFENG_IMAGE = Path(
    r"E:\量子散射\Yifeng CHEN 999016959\Codes and Figures\output"
    r"\part4_wavepacket_scattering\first_resonance.png"
)
SHUFAN_DOC = Path(r"E:\量子散射\Shufan ZHANG 999018435\quantum scattering report.docx")
SHUFAN_IMAGE = Path(
    r"E:\量子散射\Shufan ZHANG 999018435\codes and figures\outputs"
    r"\part4\part4_first_resonance.png"
)


def find_paragraph(document: Document, exact_text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == exact_text]
    if len(matches) != 1:
        raise RuntimeError(f"expected exactly one paragraph {exact_text!r}, found {len(matches)}")
    return matches[0]


def copy_paragraph_properties(source, target) -> None:
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))


def copy_run_properties(source_run, target_run) -> None:
    if source_run is None or source_run._r.rPr is None:
        return
    if target_run._r.rPr is not None:
        target_run._r.remove(target_run._r.rPr)
    target_run._r.insert(0, deepcopy(source_run._r.rPr))


def first_text_run(paragraph):
    for run in paragraph.runs:
        if run.text:
            return run
    return paragraph.runs[0] if paragraph.runs else None


def insert_text_paragraph(anchor, template, text: str):
    paragraph = anchor.insert_paragraph_before()
    copy_paragraph_properties(template, paragraph)
    run = paragraph.add_run(text)
    copy_run_properties(first_text_run(template), run)
    return paragraph


def insert_picture_paragraph(anchor, template, image_path: Path, width_emu: int):
    paragraph = anchor.insert_paragraph_before()
    copy_paragraph_properties(template, paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Emu(width_emu))
    return paragraph


def replace_text_preserving_format(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def assert_no_part5(document: Document) -> None:
    if any("part 5" in paragraph.text.lower() for paragraph in document.paragraphs):
        raise RuntimeError("Part 5 content unexpectedly present")


def edit_yifeng() -> None:
    document = Document(YIFENG_DOC)
    assert len(document.paragraphs) == 85
    assert len(document.inline_shapes) == 14
    assert_no_part5(document)

    anchor = find_paragraph(document, "4.4 Narrow packet around the second resonance")
    old_caption = find_paragraph(
        document,
        "Figure 14. Narrow packet centered at the Part 3 second peak, E₀ = 1.32882395.",
    )
    anchor_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is anchor._p
    )
    image_template = document.paragraphs[anchor_index + 1]
    caption_template = document.paragraphs[anchor_index + 2]
    body_template = document.paragraphs[anchor_index + 3]

    heading = anchor.insert_paragraph_before(
        "4.4 Narrow packet around the first resonance", style=anchor.style
    )
    copy_paragraph_properties(anchor, heading)
    insert_picture_paragraph(anchor, image_template, YIFENG_IMAGE, 4_297_680)
    insert_text_paragraph(
        anchor,
        caption_template,
        "Figure 14. Narrow packet centered at the Part 3 first peak, E₀ = 0.62097030.",
    )
    insert_text_paragraph(
        anchor,
        body_template,
        "The first resonance is much narrower than the second, so this calculation uses an "
        "extremely small momentum width. Nearly all spectral components remain inside the "
        "high-transmission peak, giving an integrated transmission probability of about "
        "0.9984. The narrow momentum distribution makes the packet extremely broad in real "
        "space, so the global panels show its envelope.",
    )

    replace_text_preserving_format(anchor, "4.5 Narrow packet around the second resonance")
    replace_text_preserving_format(
        old_caption,
        "Figure 15. Narrow packet centered at the Part 3 second peak, E₀ = 1.32882395.",
    )
    document.save(YIFENG_DOC)

    check = Document(YIFENG_DOC)
    assert len(check.paragraphs) == 89
    assert len(check.inline_shapes) == 15
    assert_no_part5(check)
    assert find_paragraph(check, "4.4 Narrow packet around the first resonance")
    assert find_paragraph(check, "4.5 Narrow packet around the second resonance")


def edit_shufan() -> None:
    document = Document(SHUFAN_DOC)
    assert len(document.paragraphs) == 81
    assert len(document.inline_shapes) == 13
    assert_no_part5(document)

    anchor = find_paragraph(document, "4.4 Narrow range at second resonance peak")
    old_caption = find_paragraph(
        document,
        "Figure 12. Extremely narrow-band packet centered exactly at the numerically determined "
        "second resonance. Its large spatial width is the Fourier counterpart of the tiny "
        "momentum width.",
    )
    anchor_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is anchor._p
    )
    body_template = document.paragraphs[anchor_index + 1]
    image_template = document.paragraphs[anchor_index + 2]
    caption_template = document.paragraphs[anchor_index + 3]

    heading = anchor.insert_paragraph_before(
        "4.4 Narrow range at first resonance peak", style=anchor.style
    )
    copy_paragraph_properties(anchor, heading)
    insert_text_paragraph(
        anchor,
        body_template,
        "The first resonance at E₀=0.62097030 is considerably narrower than the second "
        "resonance, so a₀=2.12×10⁻⁶ is used to confine almost the whole incident spectrum "
        "to the high-transmission window. The momentum-space calculation gives "
        "P_T=0.99836 and P_R=0.00164; therefore, the packet is transmitted almost "
        "completely. This very small momentum width produces a real-space packet on the "
        "scale of 10⁶, so the global snapshots display the ±|Ψ| envelope, while the local "
        "animation resolves the rapidly oscillating complex carrier near the potential.",
    )
    insert_picture_paragraph(anchor, image_template, SHUFAN_IMAGE, 5_897_880)
    insert_text_paragraph(
        anchor,
        caption_template,
        "Figure 12. First-resonance packet centered at E₀=0.62097030. The global panels show "
        "the broad magnitude envelope required by its extremely small momentum width.",
    )

    replace_text_preserving_format(anchor, "4.5 Narrow range at second resonance peak")
    replace_text_preserving_format(
        old_caption,
        "Figure 13. Extremely narrow-band packet centered exactly at the numerically determined "
        "second resonance. Its large spatial width is the Fourier counterpart of the tiny "
        "momentum width.",
    )
    document.save(SHUFAN_DOC)

    check = Document(SHUFAN_DOC)
    assert len(check.paragraphs) == 85
    assert len(check.inline_shapes) == 14
    assert_no_part5(check)
    assert find_paragraph(check, "4.4 Narrow range at first resonance peak")
    assert find_paragraph(check, "4.5 Narrow range at second resonance peak")


def main() -> None:
    for path in (YIFENG_DOC, YIFENG_IMAGE, SHUFAN_DOC, SHUFAN_IMAGE):
        if not path.exists():
            raise FileNotFoundError(path)
    edit_yifeng()
    edit_shufan()
    print("Edited both reports successfully.")


if __name__ == "__main__":
    main()

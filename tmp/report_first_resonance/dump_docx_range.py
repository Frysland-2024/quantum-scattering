from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def main() -> None:
    if len(sys.argv) not in (2, 4):
        raise SystemExit("usage: dump_docx_range.py DOCX [START END]")
    path = Path(sys.argv[1])
    document = Document(path)
    start = int(sys.argv[2]) if len(sys.argv) == 4 else 0
    end = int(sys.argv[3]) if len(sys.argv) == 4 else len(document.paragraphs)
    for index, paragraph in enumerate(document.paragraphs[start:end], start=start):
        drawings = len(paragraph._p.xpath(".//w:drawing"))
        print(
            f"{index:03d}\t{paragraph.style.name}\t{paragraph.alignment}\t"
            f"drawings={drawings}\t{paragraph.text}"
        )


if __name__ == "__main__":
    main()

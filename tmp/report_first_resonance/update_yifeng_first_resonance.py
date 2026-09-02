from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image


SOURCE = Path(
    r"E:\量子散射\Yifeng CHEN 999016959\Quantum Scattering Project Report.docx"
)
FIGURE = Path(
    r"E:\量子散射\Yifeng CHEN 999016959\Codes and Figures\output"
    r"\part4_wavepacket_scattering\first_resonance.png"
)
OUTPUT = Path(
    r"E:\量子散射\tmp\report_first_resonance\Yifeng_report_teacher_gif_format.docx"
)

OLD_SENTENCE = (
    "The narrow momentum distribution makes the packet extremely broad in real "
    "space, so the global panels show its envelope."
)
NEW_SENTENCE = (
    "The narrow momentum distribution makes the packet extremely broad in real "
    "space. The full-range panels therefore show the real and imaginary parts over "
    "the complete propagation distance, while the local animation resolves the "
    "oscillations close to the barrier."
)


def first_paragraph_index(document: Document, predicate) -> int:
    matches = [index for index, paragraph in enumerate(document.paragraphs) if predicate(paragraph)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one matching paragraph, found {matches}")
    return matches[0]


def main() -> None:
    document = Document(SOURCE)
    heading_index = first_paragraph_index(
        document,
        lambda paragraph: paragraph.text.strip() == "4.4 Narrow packet around the first resonance",
    )
    caption_index = first_paragraph_index(
        document,
        lambda paragraph: paragraph.text.startswith("Figure 14. Narrow packet centered"),
    )
    body_index = caption_index + 1
    body = document.paragraphs[body_index]
    if OLD_SENTENCE not in body.text:
        raise RuntimeError("The expected first-resonance report sentence was not found.")
    if len(body.runs) != 1:
        raise RuntimeError(f"Expected one body run, found {len(body.runs)}")
    body.runs[0].text = body.text.replace(OLD_SENTENCE, NEW_SENTENCE)

    figure_paragraphs = [
        paragraph
        for paragraph in document.paragraphs[heading_index + 1 : caption_index]
        if paragraph._p.xpath(".//w:drawing")
    ]
    if len(figure_paragraphs) != 1:
        raise RuntimeError(f"Expected one Figure 14 drawing paragraph, found {len(figure_paragraphs)}")
    figure_paragraph = figure_paragraphs[0]
    blips = figure_paragraph._p.xpath(".//a:blip")
    inlines = figure_paragraph._p.xpath(".//wp:inline")
    if len(blips) != 1 or len(inlines) != 1:
        raise RuntimeError("Figure 14 is not represented by one inline image.")

    relationship_id = blips[0].get(qn("r:embed"))
    image_part = document.part.rels[relationship_id].target_part
    figure_bytes = FIGURE.read_bytes()
    image_part._blob = figure_bytes
    image_part._image = None

    with Image.open(FIGURE) as image:
        pixel_width, pixel_height = image.size
    inline = inlines[0]
    extent = inline.xpath("./wp:extent")[0]
    transform_extent = inline.xpath(".//pic:spPr/a:xfrm/a:ext")[0]
    width_emu = int(extent.get("cx"))
    height_emu = round(width_emu * pixel_height / pixel_width)
    extent.set("cy", str(height_emu))
    transform_extent.set("cx", str(width_emu))
    transform_extent.set("cy", str(height_emu))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    check = Document(OUTPUT)
    check_body_index = first_paragraph_index(
        check,
        lambda paragraph: paragraph.text.startswith(
            "The first resonance is much narrower than the second"
        ),
    )
    if NEW_SENTENCE not in check.paragraphs[check_body_index].text:
        raise RuntimeError("The revised explanation did not survive the DOCX round trip.")
    check_caption_index = first_paragraph_index(
        check,
        lambda paragraph: paragraph.text.startswith("Figure 14. Narrow packet centered"),
    )
    check_figure = next(
        paragraph
        for paragraph in reversed(check.paragraphs[:check_caption_index])
        if paragraph._p.xpath(".//w:drawing")
    )
    check_blip = check_figure._p.xpath(".//a:blip")[0]
    check_relationship_id = check_blip.get(qn("r:embed"))
    if check.part.rels[check_relationship_id].target_part.blob != figure_bytes:
        raise RuntimeError("The revised Figure 14 image bytes do not match the generated PNG.")
    print(OUTPUT)
    print(f"Figure 14 extent: {width_emu} x {height_emu} EMU")


if __name__ == "__main__":
    main()

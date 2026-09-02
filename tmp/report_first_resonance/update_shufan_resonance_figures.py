from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu


SOURCE = Path(r"E:\量子散射\Shufan ZHANG 999018435\quantum scattering report.docx")
TARGET = Path(
    r"E:\量子散射\tmp\report_first_resonance"
    r"\Shufan_report_teacher_output_sync.docx"
)
FIRST_IMAGE = Path(
    r"E:\量子散射\Shufan ZHANG 999018435\codes and figures\outputs"
    r"\part4\part4_first_resonance.png"
)
SECOND_IMAGE = Path(
    r"E:\量子散射\Shufan ZHANG 999018435\codes and figures\outputs"
    r"\part4\part4_second_resonance.png"
)
FIGURE_WIDTH = 5_897_880


def find_exact(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"expected one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def replace_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_picture(paragraph, image_path: Path) -> None:
    if len(paragraph._p.xpath(".//w:drawing")) != 1:
        raise RuntimeError("expected exactly one drawing in target paragraph")
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run().add_picture(str(image_path), width=Emu(FIGURE_WIDTH))


def main() -> None:
    for path in (SOURCE, FIRST_IMAGE, SECOND_IMAGE):
        if not path.exists():
            raise FileNotFoundError(path)

    document = Document(SOURCE)
    if len(document.paragraphs) != 85 or len(document.inline_shapes) != 14:
        raise RuntimeError("unexpected Shufan report structure")

    first_body = find_exact(
        document,
        "The first resonance at E₀=0.62097030 is considerably narrower than the second "
        "resonance, so a₀=2.12×10⁻⁶ is used to confine almost the whole incident spectrum "
        "to the high-transmission window. The momentum-space calculation gives "
        "P_T=0.99836 and P_R=0.00164; therefore, the packet is transmitted almost "
        "completely. This very small momentum width produces a real-space packet on the "
        "scale of 10⁶, so the global snapshots display the ±|Ψ| envelope, while the local "
        "animation resolves the rapidly oscillating complex carrier near the potential.",
    )
    first_caption = find_exact(
        document,
        "Figure 12. First-resonance packet centered at E₀=0.62097030. The global panels "
        "show the broad magnitude envelope required by its extremely small momentum width.",
    )
    second_caption = find_exact(
        document,
        "Figure 13. Extremely narrow-band packet centered exactly at the numerically "
        "determined second resonance. Its large spatial width is the Fourier counterpart "
        "of the tiny momentum width.",
    )

    paragraphs = document.paragraphs
    first_caption_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph._p is first_caption._p
    )
    second_caption_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph._p is second_caption._p
    )
    first_picture = paragraphs[first_caption_index - 1]
    second_picture = paragraphs[second_caption_index - 1]

    replace_text(
        first_body,
        "The first resonance at E₀=0.62097030 is considerably narrower than the second "
        "resonance, so a₀=2.12×10⁻⁶ is used to confine almost the whole incident spectrum "
        "to the high-transmission window. The momentum-space calculation gives "
        "P_T=0.99836 and P_R=0.00164; therefore, the packet is transmitted almost "
        "completely. This very small momentum width produces a real-space packet on the "
        "scale of 10⁶. In the full-range panels, the red real part and dashed blue imaginary "
        "part are retained over the complete propagation distance; the local animation then "
        "shows their rapid carrier oscillations around the potential.",
    )
    replace_text(
        first_caption,
        "Figure 12. First-resonance packet centered at E₀=0.62097030. The three full-range "
        "snapshots show the real and imaginary parts before, during and after transmission.",
    )
    replace_picture(first_picture, FIRST_IMAGE)
    replace_picture(second_picture, SECOND_IMAGE)

    document.save(TARGET)

    check = Document(TARGET)
    if len(check.paragraphs) != 85 or len(check.inline_shapes) != 14:
        raise RuntimeError("report structure changed unexpectedly")
    if any("±|Ψ| envelope" in paragraph.text for paragraph in check.paragraphs):
        raise RuntimeError("stale envelope wording remains")
    print(TARGET)


if __name__ == "__main__":
    main()

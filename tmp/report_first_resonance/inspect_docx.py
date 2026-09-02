from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def fmt(value):
    if value is None:
        return "-"
    return str(value)


def inspect(path: Path) -> None:
    doc = Document(path)
    print(f"FILE\t{path}")
    print(f"PARAGRAPHS\t{len(doc.paragraphs)}")
    print(f"TABLES\t{len(doc.tables)}")
    print(f"INLINE_SHAPES\t{len(doc.inline_shapes)}")
    print(f"SECTIONS\t{len(doc.sections)}")
    print("INLINE_SHAPE_INDEX\tWIDTH_EMU\tHEIGHT_EMU")
    for shape_idx, shape in enumerate(doc.inline_shapes):
        print(f"{shape_idx}\t{shape.width}\t{shape.height}")
    print("PARAGRAPH_INDEX\tSTYLE\tALIGN\tBEFORE\tAFTER\tLINE\tDRAWINGS\tPAGE_BREAK\tTEXT")
    for idx, para in enumerate(doc.paragraphs):
        drawings = para._p.xpath(".//w:drawing")
        page_break = bool(para._p.xpath(".//w:br[@w:type='page']"))
        text = para.text.replace("\t", " ").replace("\n", "\\n")
        if text.strip() or drawings or page_break:
            pf = para.paragraph_format
            print(
                "\t".join(
                    [
                        str(idx),
                        para.style.name if para.style is not None else "",
                        fmt(para.alignment),
                        fmt(pf.space_before),
                        fmt(pf.space_after),
                        fmt(pf.line_spacing),
                        str(len(drawings)),
                        str(page_break),
                        text,
                    ]
                )
            )

            for run_idx, run in enumerate(para.runs):
                if not (run.text.strip() or run._r.xpath(".//w:drawing")):
                    continue
                font = run.font
                print(
                    "\tRUN\t"
                    + "\t".join(
                        [
                            str(run_idx),
                            repr(run.text),
                            f"bold={font.bold}",
                            f"italic={font.italic}",
                            f"size={fmt(font.size)}",
                            f"name={fmt(font.name)}",
                            f"drawing={bool(run._r.xpath('.//w:drawing'))}",
                        ]
                    )
                )

    print("BLOCK_ORDER")
    p_index = {id(p._p): i for i, p in enumerate(doc.paragraphs)}
    t_index = {id(t._tbl): i for i, t in enumerate(doc.tables)}
    body = doc.element.body
    for block_idx, child in enumerate(body.iterchildren()):
        if child.tag == qn("w:p"):
            print(f"{block_idx}\tP\t{p_index.get(id(child), '?')}")
        elif child.tag == qn("w:tbl"):
            print(f"{block_idx}\tT\t{t_index.get(id(child), '?')}")
        elif child.tag == qn("w:sectPr"):
            print(f"{block_idx}\tSECTPR")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    inspect(args.docx)


if __name__ == "__main__":
    main()

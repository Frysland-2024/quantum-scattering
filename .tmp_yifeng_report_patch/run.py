from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import os, tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt
import matplotlib.pyplot as plt

REPORT = Path('Yifeng CHEN 999016959/Quantum_Scattering_Project_Report.docx')


def find_para(doc, exact):
    for p in doc.paragraphs:
        if p.text == exact:
            return p
    raise RuntimeError(f'Paragraph not found: {exact}')


def set_text(p, text, font='Times New Roman', size=11, italic=None, bold=None, alignment=None):
    # Clear runs while retaining paragraph properties/style.
    for r in list(p.runs):
        p._p.remove(r._r)
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    if italic is not None:
        r.italic = italic
    if bold is not None:
        r.bold = bold
    if alignment is not None:
        p.alignment = alignment
    return p


def insert_after(p, text, font='Times New Roman', size=11, alignment=None):
    new_p = OxmlElement('w:p')
    p._p.addnext(new_p)
    q = Paragraph(new_p, p._parent)
    q.style = p.style
    if alignment is not None:
        q.alignment = alignment
    r = q.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    return q


def make_table_png(path):
    rows = [
        ['10', '0.00121665', '0.00121931', '0.00000266'],
        ['100', '0.00121908', '0.00121931', '0.00000023'],
        ['1000', '0.00121935', '0.00121931', '0.00000004'],
    ]
    headers = ['L', r'$\int_{-10}^{10} f(p)\,\delta_L(p)\,dp$', r'$f(0)$', 'Absolute difference']
    fig, ax = plt.subplots(figsize=(12, 4.35), dpi=100)
    ax.axis('off')
    ax.set_title('Example of a test calculation', fontsize=28, fontweight='bold', pad=22)
    tab = ax.table(cellText=rows, colLabels=headers, cellLoc='center', colLoc='center', loc='center')
    tab.auto_set_font_size(False)
    tab.set_fontsize(16)
    tab.scale(1.0, 2.55)
    for cell in tab.get_celld().values():
        cell.set_linewidth(1.0)
    fig.tight_layout(pad=0.4)
    fig.savefig(path, dpi=100)
    plt.close(fig)


doc = Document(REPORT)

# ---- Part 1: derive the free plane-wave basis before the Fourier expansion ----
p0 = find_para(doc, 'ψ(0,x) = (2a/π)¹ᐟ⁴ exp[−a(x−x₀)²] exp[i p₀(x−x₀)]')
q = insert_after(p0,
    'Using m = ℏ = 1, the free-particle time-independent Schrödinger equation is',
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
q = insert_after(q, '−(1/2) d²uₚ(x)/dx² = E uₚ(x).',
    font='Cambria Math', alignment=WD_ALIGN_PARAGRAPH.CENTER)
q = insert_after(q,
    'With E = p²/2, this becomes d²uₚ(x)/dx² + p²uₚ(x) = 0, whose solutions are plane waves exp(±ipx). For the momentum eigenstate with momentum p, the plane-wave form used in the Fourier expansion is',
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
q = insert_after(q, 'uₚ(x) = 1/√(2π) exp(ipx).',
    font='Cambria Math', alignment=WD_ALIGN_PARAGRAPH.CENTER)

set_text(find_para(doc, 'The propagation is done in momentum space. The initial packet is Fourier transformed as'),
    'The initial Gaussian packet can therefore be expanded in these free-particle momentum eigenstates. Its momentum-space amplitude Φ(p) is obtained from the Fourier transform',
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
set_text(find_para(doc, 'Each momentum component then gains the free-particle phase exp(−ip²t/2), giving'),
    'During free propagation, each momentum component keeps the same amplitude Φ(p) and gains the free-particle phase exp(−ip²t/2), giving',
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

# ---- Part 2: define delta first, then introduce the finite-L regularization ----
p = find_para(doc, 'The approximation function of Dirac delta function used in the project is')
set_text(p, 'The Dirac delta is defined by its action on a smooth test function f(p):',
         alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
p = find_para(doc, 'δL(p) = sin(Lp)/(πp).')
set_text(p, '∫₋∞⁺∞ f(p) δ(p) dp = f(0).', font='Cambria Math', alignment=WD_ALIGN_PARAGRAPH.CENTER)
q = insert_after(p, 'A Fourier representation of the same distribution is',
                 alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
q = insert_after(q, 'δ(p) = 1/(2π) ∫₋∞⁺∞ exp(ipx) dx.',
                 font='Cambria Math', alignment=WD_ALIGN_PARAGRAPH.CENTER)
q = insert_after(q, 'For the numerical calculation, the infinite x-range is truncated to [−L,L]. This gives the regularized delta function',
                 alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
q = insert_after(q, 'δ_L(p) = 1/(2π) ∫₋Lᴸ exp(ipx) dx = sin(Lp)/(πp).',
                 font='Cambria Math', alignment=WD_ALIGN_PARAGRAPH.CENTER)
q = insert_after(q, 'At p = 0, the limiting value is δ_L(0) = L/π. As L → ∞, δ_L(p) approaches δ(p) in the distributional sense, so ∫ f(p)δ_L(p) dp → f(0).',
                 alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

set_text(find_para(doc, 'Figure 5. Numerical test of ∫ φ(p)δL(p) dp on the fixed interval [−10,10].'),
         'Figure 5. Numerical test of ∫ f(p)δ_L(p) dp on the fixed interval [−10,10].',
         size=8, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
old_target = next(p for p in doc.paragraphs if p.text.startswith('The target value is φ(0) = 0.0012193111.'))
set_text(old_target,
         old_target.text.replace('φ(0)', 'f(0)'),
         alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

# Save text/layout edits first.
fd, tmp_docx_name = tempfile.mkstemp(suffix='.docx', dir=str(REPORT.parent))
os.close(fd)
tmp_docx = Path(tmp_docx_name)
doc.save(tmp_docx)

# Replace Figure 5's embedded table so its notation also uses f(p), not φ(p).
fd, png_name = tempfile.mkstemp(suffix='.png')
os.close(fd)
png = Path(png_name)
make_table_png(png)
fd, out_name = tempfile.mkstemp(suffix='.docx', dir=str(REPORT.parent))
os.close(fd)
out = Path(out_name)
try:
    with ZipFile(tmp_docx, 'r') as zin, ZipFile(out, 'w', ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = png.read_bytes() if item.filename == 'word/media/image5.png' else zin.read(item.filename)
            zout.writestr(item, data)
    out.replace(REPORT)
finally:
    for f in (tmp_docx, png, out):
        if f.exists():
            f.unlink()

# Integrity / scope checks.
check = Document(REPORT)
texts = [p.text for p in check.paragraphs]
assert 'uₚ(x) = 1/√(2π) exp(ipx).' in texts
assert '∫₋∞⁺∞ f(p) δ(p) dp = f(0).' in texts
assert any(t.startswith('Part 3 — Stationary scattering') for t in texts)
print('Yifeng report patched successfully.')

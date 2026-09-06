from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

src = Path('Yifeng CHEN 999016959/Quantum_Scattering_Project_Report.docx')
doc = Document(src)

def find_para(text):
    for p in doc.paragraphs:
        if p.text == text:
            return p
    raise ValueError(text)

def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn('w:pPr'):
            p._p.remove(child)

def remove_paragraph(p):
    p._element.getparent().remove(p._element)
    p._p = p._element = None

def set_body(p, text):
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

M = 'm:'
def m_el(tag, val=None):
    e = OxmlElement(M + tag)
    if val is not None:
        e.set(qn('m:val'), val)
    return e

def math_run(text, italic=True):
    r = m_el('r')
    rpr = m_el('rPr')
    rpr.append(m_el('sty', 'i' if italic else 'p'))
    r.append(rpr)
    t = m_el('t'); t.text = text; r.append(t)
    return r

def add_text(parent, text, italic=True):
    parent.append(math_run(text, italic))

def add_subscript(parent, base, sub):
    ss = m_el('sSub')
    e = m_el('e'); add_text(e, base, True)
    s = m_el('sub'); add_text(s, sub, True)
    ss.append(e); ss.append(s); parent.append(ss)

def set_equation_integral(p):
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    omp = m_el('oMathPara'); om = m_el('oMath')
    nary = m_el('nary')
    npr = m_el('naryPr'); npr.append(m_el('chr','∫')); npr.append(m_el('limLoc','undOvr')); nary.append(npr)
    sub = m_el('sub'); add_text(sub,'−∞',False)
    sup = m_el('sup'); add_text(sup,'∞',False)
    expr = m_el('e')
    for txt,it in [('f',True),('(',False),('p',True),(')',False),('δ',True),('(',False),('p',True),(')',False),('d',True),('p',True)]: add_text(expr,txt,it)
    nary.extend([sub,sup,expr]); om.append(nary)
    for txt,it in [(' = ',False),('f',True),('(',False),('0',False),(')',False)]: add_text(om,txt,it)
    omp.append(om); p._p.append(omp)

def set_equation_regularized(p):
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(5)
    omp = m_el('oMathPara'); om = m_el('oMath')
    add_subscript(om,'δ','L')
    for txt,it in [('(',False),('p',True),(') = ',False)]: add_text(om,txt,it)
    frac = m_el('f'); num = m_el('num'); den = m_el('den')
    func = m_el('func'); func.append(m_el('funcPr'))
    fname = m_el('fName'); add_text(fname,'sin',False)
    e = m_el('e')
    for txt,it in [('(',False),('L',True),('p',True),(')',False)]: add_text(e,txt,it)
    func.extend([fname,e]); num.append(func)
    add_text(den,'π',True); add_text(den,'p',True)
    frac.extend([num,den]); om.append(frac); omp.append(om); p._p.append(omp)

p32=find_para('The Dirac delta is defined by its action on a smooth test function f(p):')
p33=find_para('∫₋∞⁺∞ f(p) δ(p) dp = f(0).')
p34=find_para('A Fourier representation of the same distribution is')
p35=find_para('δ(p) = 1/(2π) ∫₋∞⁺∞ exp(ipx) dx.')
p36=find_para('For the numerical calculation, the infinite x-range is truncated to [−L,L]. This gives the regularized delta function')
p37=find_para('δ_L(p) = 1/(2π) ∫₋Lᴸ exp(ipx) dx = sin(Lp)/(πp).')
p38=find_para('At p = 0, the limiting value is δ_L(0) = L/π. As L → ∞, δ_L(p) approaches δ(p) in the distributional sense, so ∫ f(p)δ_L(p) dp → f(0).')

set_body(p32,'The Dirac delta function is defined by its action on a smooth test function f(p):')
set_equation_integral(p33)
set_body(p34,'Here f(p) is a smooth test function. The regularized Dirac delta function used in this project is')
set_equation_regularized(p35)
for p in (p36,p37,p38): remove_paragraph(p)

cap=find_para('Figure 5. Numerical test of ∫ f(p)δ_L(p) dp on the fixed interval [−10,10].')
for r in cap.runs:
    r.font.name='Times New Roman'; r.font.size=Pt(8)
cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

doc.save(src)
print('patched')

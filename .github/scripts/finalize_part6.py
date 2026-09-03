from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_formula(doc: Document, text: str, size: float) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(size)


def add_image(doc: Document, path: Path, width: float) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def add_shufan_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.add_run(text)


def add_yifeng_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8)


def has_part6(doc: Document) -> bool:
    return any("Part 6" in paragraph.text for paragraph in doc.paragraphs)


def append_shufan_report(report: Path, images: Path) -> None:
    doc = Document(report)
    if has_part6(doc):
        print(f"Part 6 already present: {report}")
        return

    doc.add_page_break()
    doc.add_paragraph("6. Part 6 - Complex scaling and resonance states", style="Heading 1")
    doc.add_paragraph(
        "Part 6 applies complex scaling to the same short-range potential. The coordinate is rotated as "
        "x→x exp(iθ), so the Hamiltonian becomes non-Hermitian and its eigenvalues can be complex. "
        "The calculation uses J=2000, L=200 and N=10000 for θ=0, 0.05, 0.10, 0.15, 0.20 and 0.25."
    )
    add_formula(doc, "Hθ = −(1/2) exp(−2iθ) d²/dx² + V[x exp(iθ)].", 11)

    doc.add_paragraph("6.1 Complex eigenvalue spectra", style="Heading 2")
    doc.add_paragraph(
        "At θ=0, the discretized continuum lies on the real-energy axis. As θ increases, these states rotate "
        "into the lower half of the complex-energy plane. A few isolated eigenvalues remain almost at the same "
        "positions; these are the resonance states."
    )
    add_image(doc, images / "part6_complex_eigenvalues_all.png", 6.15)
    add_shufan_caption(
        doc,
        "Figure 19. Complex-scaled eigenvalues for several values of θ. The continuum rotates downward, while isolated resonance eigenvalues remain nearly fixed.",
    )

    doc.add_page_break()
    add_image(doc, images / "part6_complex_eigenvalues_separate.png", 6.15)
    add_shufan_caption(doc, "Figure 20. The same complex-energy spectra shown separately for each value of θ.")
    doc.add_paragraph(
        "The first two stable complex energies are approximately E₁=0.620971−0.000058i and "
        "E₂=1.327197−0.015447i. Their real parts are close to the two transmission peaks found in Part 3. "
        "The much smaller imaginary part of E₁ is consistent with the first peak being much narrower than the second."
    )

    doc.add_page_break()
    doc.add_paragraph("6.2 Breit-Wigner profiles", style="Heading 2")
    doc.add_paragraph(
        "A resonance energy can be written as Eₙ=Eres,n−iΓₙ/2. Its real part gives the resonance position, "
        "while Γₙ gives the width of the corresponding Breit-Wigner peak."
    )
    add_formula(doc, "|T(E)|² ≈ (Γₙ/2)² / [(E−Eres,n)² + (Γₙ/2)²].", 11)
    add_image(doc, images / "part6_breit_wigner_all.png", 5.65)
    add_shufan_caption(
        doc,
        "Figure 21. The Hermitian transmission profile compared with Breit-Wigner profiles obtained from the selected complex resonance energies.",
    )
    add_image(doc, images / "part6_breit_wigner_zoom.png", 5.75)
    add_shufan_caption(doc, "Figure 22. Detailed comparison around the first and second transmission peaks.")
    doc.add_paragraph(
        "Around the first resonance, the Breit-Wigner curves almost overlap the Hermitian transmission curve. "
        "For the second resonance, they reproduce the peak position and width but differ away from the maximum, "
        "where the transmission also contains a non-resonant background and contributions from broader resonances."
    )

    doc.add_page_break()
    doc.add_paragraph("6.3 Resonance wavefunction", style="Heading 2")
    doc.add_paragraph(
        "The last two figures show the second resonance state at Eres≈1.327197−0.015447i. "
        "The original Siegert state satisfies outgoing boundary conditions and grows at large |x|, so it is not "
        "square integrable. After the rotation ψθ(x)=ψ[x exp(iθ)] with θ=0.1, the same state decays away from the potential."
    )
    add_image(doc, images / "part6_resonance_wavefunction_near.png", 5.85)
    add_shufan_caption(
        doc,
        "Figure 23. The second resonance wavefunction over −200≤x≤200 before and after complex scaling with θ=0.1.",
    )
    add_image(doc, images / "part6_resonance_wavefunction_far.png", 5.85)
    add_shufan_caption(
        doc,
        "Figure 24. The same resonance state over −500≤x≤500, showing the outward growth of the unscaled state and the decay of the complex-scaled state.",
    )
    doc.add_paragraph(
        "Therefore, complex scaling identifies resonance energies close to the transmission peaks from Part 3, "
        "explains their Breit-Wigner shapes, and converts the divergent resonance wavefunction into a decaying function that can be treated numerically."
    )
    doc.save(report)


def append_yifeng_report(report: Path, images: Path) -> None:
    doc = Document(report)
    if has_part6(doc):
        print(f"Part 6 already present: {report}")
        return

    doc.add_page_break()
    doc.add_paragraph("Part 6 — Complex scaling and resonance states", style="Heading 1")
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(
        "The short-range potential is now calculated with the complex-scaled coordinate x→x exp(iθ). "
        "This makes the Hamiltonian non-Hermitian, so resonance states can appear as complex eigenvalues. "
        "The calculation uses J=2000, L=200, N=10000 and θ=0, 0.05, 0.10, 0.15, 0.20, and 0.25."
    )
    add_formula(doc, "Hθ = −(1/2) exp(−2iθ) d²/dx² + V[x exp(iθ)].", 10)

    doc.add_paragraph("6.1 Complex eigenvalue spectra", style="Heading 2")
    add_image(doc, images / "complex_eigenvalues_all.png", 6.0)
    add_yifeng_caption(doc, "Figure 21. Complex eigenvalues for several values of θ shown in one plot.")
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(
        "At θ=0, the box continuum lies on the real axis. As θ increases, these states rotate downward into the "
        "complex-energy plane. The isolated points that stay nearly fixed are the resonance eigenvalues."
    )

    doc.add_page_break()
    add_image(doc, images / "complex_eigenvalues_by_theta.png", 6.0)
    add_yifeng_caption(doc, "Figure 22. The same complex-energy spectra shown separately for each value of θ.")
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(
        "The first two resonance energies are E₁≈0.620971−0.000058i and E₂≈1.327197−0.015447i. "
        "Their real parts are close to the two transmission peaks from Part 3. The first resonance has a much "
        "smaller imaginary part, which agrees with its much narrower transmission peak."
    )

    doc.add_page_break()
    doc.add_paragraph("6.2 Breit-Wigner profiles", style="Heading 2")
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(
        "For a complex resonance energy Eₙ=Eres,n−iΓₙ/2, the real part gives the peak position and Γₙ gives its width. "
        "The corresponding transmission peak has the Breit-Wigner form"
    )
    add_formula(doc, "|T(E)|² ≈ (Γₙ/2)² / [(E−Eres,n)² + (Γₙ/2)²].", 10)
    add_image(doc, images / "breit_wigner_profiles_all.png", 5.75)
    add_yifeng_caption(doc, "Figure 23. Hermitian transmission and Breit-Wigner profiles from the complex resonance energies.")
    add_image(doc, images / "breit_wigner_profiles_zoom.png", 5.8)
    add_yifeng_caption(doc, "Figure 24. Detailed comparison around the first and second resonance peaks.")
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(
        "The first Breit-Wigner profile almost overlaps the Hermitian result. The second one gives the peak position "
        "and width, but the curves differ away from the maximum because the transmission also contains background "
        "and broader-resonance contributions."
    )

    doc.add_page_break()
    doc.add_paragraph("6.3 Resonance wave function", style="Heading 2")
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(
        "Figures 25 and 26 show the second resonance state, Eres≈1.327197−0.015447i. The original Siegert wave "
        "function has outgoing boundary conditions and grows at large |x|. After complex scaling, "
        "ψθ(x)=ψ[x exp(iθ)] with θ=0.1 decays away from the potential."
    )
    add_image(doc, images / "resonance_wavefunction_near.png", 5.85)
    add_yifeng_caption(doc, "Figure 25. Second resonance wave function for −200≤x≤200 before and after complex scaling.")
    add_image(doc, images / "resonance_wavefunction_far.png", 5.85)
    add_yifeng_caption(doc, "Figure 26. The same resonance wave function over −500≤x≤500.")
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(
        "The larger range makes the difference clearer: the unscaled state grows outward, while the complex-scaled "
        "state remains localized. Thus the complex resonance energies explain the transmission peaks found by the "
        "Hermitian calculation and provide a direct picture of the resonance states."
    )
    doc.save(report)


def patch_yifeng_code(root: Path) -> None:
    source_path = root / "Yifeng CHEN 999016959" / "Codes and Figures" / "src" / "quantum_scattering" / "complex_scaling.py"
    text = source_path.read_text(encoding="utf-8")

    text = text.replace("import json\n", "")
    text = text.replace("from .potential import double_barrier_potential\n", "")

    reference_start = text.find("TEACHER_POLE_REFERENCES: tuple[complex, ...] = (")
    if reference_start >= 0:
        reference_end = text.find(")\n\n", reference_start)
        if reference_end < 0:
            raise RuntimeError("could not locate the end of TEACHER_POLE_REFERENCES")
        text = text[:reference_start] + text[reference_end + 3 :]

    fwhm_start = text.find("\ndef _fwhm(")
    generate_start = text.find("\ndef generate_part6(")
    if fwhm_start >= 0:
        if generate_start < 0 or generate_start <= fwhm_start:
            raise RuntimeError("could not locate generate_part6 after _fwhm")
        text = text[:fwhm_start] + "\n" + text[generate_start:]

    text = text.replace(
        '    """Generate the four teacher-style Part-6 figures and audit files."""',
        '    """Generate the four teacher-style Part-6 PNG figures."""',
    )

    cleanup_marker = '    for name in ("part6_diagnostics.json", "part6_manifest.json"):'
    if cleanup_marker not in text:
        output_line = '    output = Path(root) / "part6_complex_scaling"\n    output.mkdir(parents=True, exist_ok=True)\n'
        replacement = output_line + (
            '    for name in ("part6_diagnostics.json", "part6_manifest.json"):\n'
            '        (output / name).unlink(missing_ok=True)\n'
        )
        if output_line not in text:
            raise RuntimeError("could not locate Part 6 output directory creation")
        text = text.replace(output_line, replacement, 1)

    diagnostics_start = text.find("\n    continuum_checks: list[dict[str, object]] = []")
    return_position = text.find("\n    return created", diagnostics_start)
    if diagnostics_start >= 0:
        if return_position < 0:
            raise RuntimeError("could not locate return created after diagnostics")
        text = text[:diagnostics_start] + text[return_position:]

    text = text.replace('    "TEACHER_POLE_REFERENCES",\n', "")
    source_path.write_text(text, encoding="utf-8")

    wrapper_path = root / "Yifeng CHEN 999016959" / "Codes and Figures" / "part6.py"
    wrapper_path.write_text(
        '''"""Part 6: generate complex-scaling, Breit--Wigner, and wavefunction figures."""

from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

SRC_DIR = Path(__file__).resolve().parent / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from quantum_scattering.complex_scaling import generate_part6
from quantum_scattering.resonance_wavefunction import generate_resonance_wavefunctions


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate the six Part 6 complex-scaling and resonance figures."
    )
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    parser.add_argument(
        "--profile",
        choices=("quick", "reference", "full"),
        default="full",
        help="quick uses J=400; reference/full use the teacher's J=2000, L=200, N=10000.",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    started = time.perf_counter()
    created = generate_part6(args.output_dir, profile=args.profile)
    created.extend(
        generate_resonance_wavefunctions(args.output_dir, profile=args.profile)
    )
    figures = [
        Path(path)
        for path in created
        if Path(path).suffix.lower() == ".png" and Path(path).exists()
    ]
    elapsed = time.perf_counter() - started
    print(f"Part 6 generated {len(figures)} figures in {elapsed:.2f} s")
    for path in figures:
        print(path.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
''',
        encoding="utf-8",
    )

    output = root / "Yifeng CHEN 999016959" / "Codes and Figures" / "output" / "part6_complex_scaling"
    for name in ("part6_diagnostics.json", "part6_manifest.json"):
        (output / name).unlink(missing_ok=True)

    cleaned = source_path.read_text(encoding="utf-8")
    if "json.dumps" in cleaned or "import json" in cleaned:
        raise RuntimeError("Yifeng complex_scaling.py still contains JSON generation code")
    if "return created" not in cleaned:
        raise RuntimeError("Yifeng complex_scaling.py no longer returns its PNG list")


def validate_report(path: Path, heading: str, expected_images: int) -> None:
    doc = Document(path)
    if not any(heading in paragraph.text for paragraph in doc.paragraphs):
        raise RuntimeError(f"missing Part 6 heading in {path}")
    if len(doc.inline_shapes) != expected_images:
        raise RuntimeError(
            f"unexpected image count in {path}: {len(doc.inline_shapes)} != {expected_images}"
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[2])
    args = parser.parse_args()
    root = args.root.resolve()

    shufan_root = root / "Shufan ZHANG 999018435"
    yifeng_root = root / "Yifeng CHEN 999016959"
    append_shufan_report(
        shufan_root / "quantum_scattering_report.docx",
        shufan_root / "codes and figures" / "outputs" / "part6",
    )
    append_yifeng_report(
        yifeng_root / "Quantum_Scattering_Project_Report.docx",
        yifeng_root / "Codes and Figures" / "output" / "part6_complex_scaling",
    )
    patch_yifeng_code(root)

    validate_report(
        shufan_root / "quantum_scattering_report.docx",
        "6. Part 6 - Complex scaling and resonance states",
        25,
    )
    validate_report(
        yifeng_root / "Quantum_Scattering_Project_Report.docx",
        "Part 6 — Complex scaling and resonance states",
        26,
    )
    print("Part 6 reports and Yifeng code finalized successfully.")


if __name__ == "__main__":
    main()

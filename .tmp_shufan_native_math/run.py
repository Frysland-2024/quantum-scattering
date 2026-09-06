from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from pathlib import Path

SRC = Path('Shufan ZHANG 999018435/quantum_scattering_report.docx')
OUT = SRC
doc = Document(SRC)

M='m:'

def me(tag,val=None):
    e=OxmlElement(M+tag)
    if val is not None:
        e.set(qn('m:val'),str(val))
    return e

def clear_para(p):
    for c in list(p._p):
        if c.tag != qn('w:pPr'):
            p._p.remove(c)

def set_display(p,builder,after_pt=3):
    clear_para(p)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.space_after=Pt(after_pt)
    omp=me('oMathPara'); om=me('oMath'); builder(om); omp.append(om); p._p.append(omp)

def run(text,italic=True):
    r=me('r'); rp=me('rPr'); rp.append(me('sty','i' if italic else 'p')); r.append(rp)
    t=me('t'); t.text=text; r.append(t); return r

def t(parent,text,italic=True): parent.append(run(text,italic))
def tnor(parent,text):
    r=me('r'); rp=me('rPr'); rp.append(me('nor','1')); r.append(rp); tt=me('t'); tt.text=text; r.append(tt); parent.append(r)
def plain(text): return lambda p:t(p,text,False)
def var(text): return lambda p:t(p,text,True)

def sub(parent,base,subb):
    z=me('sSub'); e=me('e'); base(e); s=me('sub'); subb(s); z.extend([e,s]); parent.append(z)

def sup(parent,base,supp):
    z=me('sSup'); e=me('e'); base(e); s=me('sup'); supp(s); z.extend([e,s]); parent.append(z)

def subsup(parent,base,subb,supp):
    z=me('sSubSup'); e=me('e'); base(e); s1=me('sub'); subb(s1); s2=me('sup'); supp(s2); z.extend([e,s1,s2]); parent.append(z)

def frac(parent,num,den):
    z=me('f'); n=me('num'); num(n); d=me('den'); den(d); z.extend([n,d]); parent.append(z)

def rad(parent,expr):
    z=me('rad'); pr=me('radPr'); pr.append(me('degHide','1')); z.append(pr); z.append(me('deg')); e=me('e'); expr(e); z.append(e); parent.append(z)

def func(parent,name,arg):
    z=me('func'); z.append(me('funcPr')); fn=me('fName'); t(fn,name,False); e=me('e'); arg(e); z.extend([fn,e]); parent.append(z)

def delim(parent,beg,end,inner):
    z=me('d'); pr=me('dPr'); pr.append(me('begChr',beg)); pr.append(me('endChr',end)); z.append(pr); e=me('e'); inner(e); z.append(e); parent.append(z)

def nary(parent,symbol,expr,lo=None,hi=None):
    z=me('nary'); pr=me('naryPr'); pr.append(me('chr',symbol)); pr.append(me('limLoc','undOvr'))
    if lo is None: pr.append(me('subHide','1'))
    if hi is None: pr.append(me('supHide','1'))
    z.append(pr)
    sl=me('sub')
    if lo is not None: lo(sl)
    su=me('sup')
    if hi is not None: hi(su)
    e=me('e'); expr(e)
    z.extend([sl,su,e]); parent.append(z)

def acc(parent,base,chr_):
    z=me('acc'); pr=me('accPr'); pr.append(me('chr',chr_)); z.append(pr); e=me('e'); base(e); z.append(e); parent.append(z)

def bar(parent,base):
    z=me('bar'); pr=me('barPr'); pr.append(me('pos','top')); z.append(pr); e=me('e'); base(e); z.append(e); parent.append(z)

def lim_low(parent,base_text,sub_builder):
    z=me('limLow'); z.append(me('limLowPr'))
    e=me('e'); tnor(e,base_text)
    lm=me('lim'); sub_builder(lm)
    z.extend([e,lm]); parent.append(z)

def find(text):
    target=' '.join(text.replace('\n',' ').split())
    for p in doc.paragraphs:
        cur=' '.join(p.text.replace('\n',' ').split())
        if cur == target: return p
    raise RuntimeError(f'paragraph not found: {text}')

def insert_after(p,text=None,style=None,align=None,font='Times New Roman',size=10.5):
    new_p=OxmlElement('w:p'); p._p.addnext(new_p); q=Paragraph(new_p,p._parent)
    q.style=style or p.style
    if align is not None: q.alignment=align
    if text is not None:
        r=q.add_run(text); r.font.name=font; r.font.size=Pt(size)
    return q

def set_body(p,text,size=10.5):
    clear_para(p); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(size)

def x0(p): sub(p,var('x'),plain('0'))
def p0(p): sub(p,var('p'),plain('0'))
def up(p): sub(p,var('u'),var('p'))
def psiE(p): sub(p,var('ψ'),var('E'))
def psibarE(p): sub(p,lambda q:bar(q,var('ψ')),var('E'))
def phij(p): sub(p,var('φ'),var('j'))
def phii(p): sub(p,var('φ'),var('i'))
def psin(p): sub(p,var('ψ'),var('n'))
def Tij(p): sub(p,var('T'),lambda q:(t(q,'i',True),t(q,'j',True)))
def Vij(p): sub(p,var('V'),lambda q:(t(q,'i',True),t(q,'j',True)))
def Hij(p): sub(p,var('H'),lambda q:(t(q,'i',True),t(q,'j',True)))
def deltaij(p): sub(p,var('δ'),lambda q:(t(q,'i',True),t(q,'j',True)))
def Htheta(p): sub(p,var('H'),var('θ'))
def En(p): sub(p,var('E'),var('n'))
def Gaman(p): sub(p,var('Γ'),var('n'))
def Eresn(p): sub(p,var('E'),lambda q:(t(q,'res',False),t(q,',',False),t(q,'n',True)))
def expf(parent,inside): func(parent,'exp',lambda e:delim(e,'(',')',inside))

p_init = find('ψ(0,x) = (2α/π)¹⁄⁴ exp[-α(x-x₀)²] exp[ip₀(x-x₀)]')
q = insert_after(p_init, 'For a free particle, V(x)=0. Using m=ħ=1, the time-independent Schrödinger equation is', size=10.5)
qeq = insert_after(q, None, align=WD_ALIGN_PARAGRAPH.CENTER)
set_display(qeq, lambda o:(t(o,'−',False),frac(o,lambda n:t(n,'1',False),lambda d:t(d,'2',False)),t(o,' ',False),frac(o,lambda n:(sup(n,var('d'),plain('2')),up(n),t(n,'(',False),t(n,'x',True),t(n,')',False)),lambda d:(t(d,'d',True),sup(d,var('x'),plain('2')))),t(o,' = ',False),t(o,'E',True),t(o,' ',False),up(o),t(o,'(',False),t(o,'x',True),t(o,').',False)))
q = insert_after(qeq, 'With E=p²/2, the equation becomes d²uₚ(x)/dx²+p²uₚ(x)=0, whose solutions are plane waves exp(±ipx). The momentum eigenstate with momentum p used in the Fourier expansion is', size=10.5)
qeq2 = insert_after(q, None, align=WD_ALIGN_PARAGRAPH.CENTER)
set_display(qeq2, lambda o:(up(o),t(o,'(',False),t(o,'x',True),t(o,') = ',False),frac(o,lambda n:t(n,'1',False),lambda d:rad(d,lambda e:(t(e,'2',False),t(e,'π',False)))),t(o,' ',False),expf(o,lambda e:(t(e,'i',True),t(e,'p',True),t(e,'x',True))),t(o,'.',False)))
set_body(find('To propagate the wavepacket, the initial state is first transformed into momentum space,'),'The initial Gaussian wavepacket can therefore be expanded in these free-particle momentum eigenstates. Its momentum-space amplitude ϕ(p) is obtained from the Fourier transform,')

p26=find('The regularization used in the project is')
p27=find('δL(p) = sin(pL)/(πp),     δL(0) = L/π.')
p28=find('As L increases, the central peak becomes taller and narrower. Although δL(p) does not converge pointwise, it approaches the Dirac delta distribution in the sense that ∫φ(p)δL(p) dp → φ(0).')
set_body(p26,'The Dirac delta function is defined through its action on a smooth test function f(p):')
set_display(p27, lambda o:(nary(o,'∫',lambda e:(t(e,'f',True),t(e,'(',False),t(e,'p',True),t(e,')',False),t(e,' ',False),t(e,'δ',True),t(e,'(',False),t(e,'p',True),t(e,')',False),t(e,' ',False),t(e,'d',True),t(e,'p',True)),lo=lambda s:t(s,'−∞',False),hi=lambda s:t(s,'+∞',False)),t(o,' = ',False),t(o,'f',True),t(o,'(',False),t(o,'0',False),t(o,').',False)))
set_body(p28,'For the numerical calculation, the regularized Dirac delta function used in this project is')
q = insert_after(p28,None,align=WD_ALIGN_PARAGRAPH.CENTER)
set_display(q, lambda o:(sub(o,var('δ'),var('L')),t(o,'(',False),t(o,'p',True),t(o,') = ',False),frac(o,lambda n:t(n,'1',False),lambda d:(t(d,'2',False),t(d,'π',False))),t(o,' ',False),nary(o,'∫',lambda e:(expf(e,lambda z:(t(z,'i',True),t(z,'p',True),t(z,'x',True))),t(e,' ',False),t(e,'d',True),t(e,'x',True)),lo=lambda s:(t(s,'−',False),t(s,'L',True)),hi=lambda s:t(s,'L',True)),t(o,' = ',False),frac(o,lambda n:func(n,'sin',lambda e:delim(e,'(',')',lambda z:(t(z,'p',True),t(z,'L',True)))),lambda d:(t(d,'π',False),t(d,'p',True))),t(o,'.',False)))
q2 = insert_after(q,'As L → ∞, the regularized function approaches the Dirac delta distribution, so that',size=10.5)
q3 = insert_after(q2,None,align=WD_ALIGN_PARAGRAPH.CENTER)
set_display(q3, lambda o:(lim_low(o,'lim',lambda s:(t(s,'L',True),t(s,' → ∞',False))),t(o,' ',False),nary(o,'∫',lambda e:(t(e,'f',True),t(e,'(',False),t(e,'p',True),t(e,') ',False),sub(e,var('δ'),var('L')),t(e,'(',False),t(e,'p',True),t(e,') ',False),t(e,'d',True),t(e,'p',True)),lo=lambda s:t(s,'−∞',False),hi=lambda s:t(s,'+∞',False)),t(o,' = ',False),t(o,'f',True),t(o,'(',False),t(o,'0',False),t(o,').',False)))
set_body(find('The test function is the momentum-space Gaussian φ(p) with α=1 and p₀=5. The numerical integration window is fixed at p∈[-10,10] for every L.'),'For the numerical test, the smooth test function f(p) is chosen as the momentum-space Gaussian ϕ(p) from Part 1, with α=1 and p₀=5. The numerical integration window is fixed at p∈[-10,10] for every L.')
set_body(find('The error decreases from 2.66×10⁻⁶ at L=10 to 3.70×10⁻⁸ at L=1000. The convergence toward φ(0)=0.0012193111 is a direct numerical verification of the defining action of the Dirac delta distribution.'),'The error decreases from 2.66×10⁻⁶ at L=10 to 3.70×10⁻⁸ at L=1000. The convergence toward f(0)=0.0012193111 is a direct numerical verification of the defining action of the Dirac delta distribution.')

if doc.tables:
    for table in doc.tables:
        texts=[[c.text.strip() for c in row.cells] for row in table.rows]
        if len(texts)>=4 and any('Absolute difference' in x for x in texts[0]):
            table.cell(0,1).text=''
            set_display(table.cell(0,1).paragraphs[0], lambda o:nary(o,'∫',lambda e:(t(e,'f',True),t(e,'(',False),t(e,'p',True),t(e,') ',False),sub(e,var('δ'),var('L')),t(e,'(',False),t(e,'p',True),t(e,') ',False),t(e,'d',True),t(e,'p',True))))
            table.cell(0,2).text=''
            set_display(table.cell(0,2).paragraphs[0], lambda o:(t(o,'f',True),t(o,'(',False),t(o,'0',False),t(o,')',False)))

set_display(find('ψ(0,x) = (2α/π)¹⁄⁴ exp[-α(x-x₀)²] exp[ip₀(x-x₀)]'), lambda o:(t(o,'ψ',True),t(o,'(',False),t(o,'0',False),t(o,',',False),t(o,'x',True),t(o,') = ',False),sup(o,lambda e:delim(e,'(',')',lambda q:frac(q,lambda n:(t(n,'2',False),t(n,'α',True)),lambda d:t(d,'π',False))),lambda ss:frac(ss,lambda n:t(n,'1',False),lambda d:t(d,'4',False))),t(o,' ',False),expf(o,lambda q:(t(q,'−',False),t(q,'α',True),sup(q,lambda r:delim(r,'(',')',lambda z:(t(z,'x',True),t(z,'−',False),x0(z))),plain('2')))),t(o,' ',False),expf(o,lambda q:(t(q,'i',True),p0(q),delim(q,'(',')',lambda r:(t(r,'x',True),t(r,'−',False),x0(r)))))))
set_display(find('ϕ(p) = 1/√(2π) ∫ dx ψ(0,x) exp(-ipx)'), lambda o:(t(o,'ϕ',True),t(o,'(',False),t(o,'p',True),t(o,') = ',False),frac(o,lambda n:t(n,'1',False),lambda d:rad(d,lambda e:(t(e,'2',False),t(e,'π',False)))),t(o,' ',False),nary(o,'∫',lambda e:(t(e,'ψ',True),t(e,'(',False),t(e,'0',False),t(e,',',False),t(e,'x',True),t(e,') ',False),expf(e,lambda q:(t(q,'−',False),t(q,'i',True),t(q,'p',True),t(q,'x',True))),t(e,' ',False),t(e,'d',True),t(e,'x',True)))))
set_display(find('ψ(t,x) = 1/√(2π) ∫ dp ϕ(p) exp(ipx) exp(-ip²t/2)'), lambda o:(t(o,'ψ',True),t(o,'(',False),t(o,'t',True),t(o,',',False),t(o,'x',True),t(o,') = ',False),frac(o,lambda n:t(n,'1',False),lambda d:rad(d,lambda e:(t(e,'2',False),t(e,'π',False)))),t(o,' ',False),nary(o,'∫',lambda e:(t(e,'ϕ',True),t(e,'(',False),t(e,'p',True),t(e,') ',False),expf(e,lambda q:(t(q,'i',True),t(q,'p',True),t(q,'x',True))),t(e,' ',False),expf(e,lambda q:(t(q,'−',False),t(q,'i',True),sup(q,var('p'),plain('2')),t(q,'t',True),t(q,'/2',False))),t(e,' ',False),t(e,'d',True),t(e,'p',True)))))

set_display(find('V(x) = (0.5x² - 0.8) exp(-0.1x²).'), lambda o:(t(o,'V',True),t(o,'(',False),t(o,'x',True),t(o,') = ',False),delim(o,'(',')',lambda e:(t(e,'0.5',False),sup(e,var('x'),plain('2')),t(e,' − 0.8',False))),t(o,' ',False),expf(o,lambda e:(t(e,'−0.1',False),sup(e,var('x'),plain('2')))),t(o,'.',False)))
p41=find('ψE(x→−∞) = exp(ikx) + R(E) exp(−ikx) ψE(x→+∞) = T(E) exp(ikx)')
set_display(p41, lambda o:(psiE(o),t(o,'(',False),t(o,'x',True),t(o,' → −∞',False),t(o,') = ',False),expf(o,lambda e:(t(e,'i',True),t(e,'k',True),t(e,'x',True))),t(o,' + ',False),t(o,'R',True),t(o,'(',False),t(o,'E',True),t(o,') ',False),expf(o,lambda e:(t(e,'−',False),t(e,'i',True),t(e,'k',True),t(e,'x',True)))))
p41b=insert_after(p41,None,align=WD_ALIGN_PARAGRAPH.CENTER)
set_display(p41b, lambda o:(psiE(o),t(o,'(',False),t(o,'x',True),t(o,' → +∞',False),t(o,') = ',False),t(o,'T',True),t(o,'(',False),t(o,'E',True),t(o,') ',False),expf(o,lambda e:(t(e,'i',True),t(e,'k',True),t(e,'x',True)))))
p43=find('ψ̄E(x→−∞) = A(E) exp(ikx) + B(E) exp(−ikx) ψ̄E(x→+∞) = exp(ikx)')
set_display(p43, lambda o:(psibarE(o),t(o,'(',False),t(o,'x',True),t(o,' → −∞',False),t(o,') = ',False),t(o,'A',True),t(o,'(',False),t(o,'E',True),t(o,') ',False),expf(o,lambda e:(t(e,'i',True),t(e,'k',True),t(e,'x',True))),t(o,' + ',False),t(o,'B',True),t(o,'(',False),t(o,'E',True),t(o,') ',False),expf(o,lambda e:(t(e,'−',False),t(e,'i',True),t(e,'k',True),t(e,'x',True)))))
p43b=insert_after(p43,None,align=WD_ALIGN_PARAGRAPH.CENTER)
set_display(p43b, lambda o:(psibarE(o),t(o,'(',False),t(o,'x',True),t(o,' → +∞',False),t(o,') = ',False),expf(o,lambda e:(t(e,'i',True),t(e,'k',True),t(e,'x',True)))))
set_display(find('ψ̄E(xₙ₋₁) = [2 + 2(Δx)²(V(xₙ)-E)] ψ̄E(xₙ) - ψ̄E(xₙ₊₁).'), lambda o:(psibarE(o),t(o,'(',False),sub(o,var('x'),lambda s:t(s,'n−1',False)),t(o,') = ',False),delim(o,'[',']',lambda e:(t(e,'2 + 2',False),sup(e,lambda q:delim(q,'(',')',lambda r:t(r,'Δx',False)),plain('2')),delim(e,'(',')',lambda q:(t(q,'V',True),t(q,'(',False),sub(q,var('x'),var('n')),t(q,') − ',False),t(q,'E',True))))),t(o,' ',False),psibarE(o),t(o,'(',False),sub(o,var('x'),var('n')),t(o,') − ',False),psibarE(o),t(o,'(',False),sub(o,var('x'),lambda s:t(s,'n+1',False)),t(o,').',False)))
set_display(find('ψE(x) = ψ̄E(x)/A(E),     T(E)=1/A(E),     R(E)=B(E)/A(E).'), lambda o:(psiE(o),t(o,'(',False),t(o,'x',True),t(o,') = ',False),frac(o,lambda n:(psibarE(n),t(n,'(',False),t(n,'x',True),t(n,')',False)),lambda d:(t(d,'A',True),t(d,'(',False),t(d,'E',True),t(d,')',False))),t(o,',      ',False),t(o,'T',True),t(o,'(',False),t(o,'E',True),t(o,') = ',False),frac(o,lambda n:t(n,'1',False),lambda d:(t(d,'A',True),t(d,'(',False),t(d,'E',True),t(d,')',False))),t(o,',      ',False),t(o,'R',True),t(o,'(',False),t(o,'E',True),t(o,') = ',False),frac(o,lambda n:(t(n,'B',True),t(n,'(',False),t(n,'E',True),t(n,')',False)),lambda d:(t(d,'A',True),t(d,'(',False),t(d,'E',True),t(d,')',False))),t(o,'.',False)))

set_display(find('Ψ(t,x) = ∫ dp ϕ(p) ψₚ(x) exp[-iE(p)t],    E(p)=p²/2.'), lambda o:(t(o,'Ψ',True),t(o,'(',False),t(o,'t',True),t(o,',',False),t(o,'x',True),t(o,') = ',False),nary(o,'∫',lambda e:(t(e,'ϕ',True),t(e,'(',False),t(e,'p',True),t(e,') ',False),sub(e,var('ψ'),var('p')),t(e,'(',False),t(e,'x',True),t(e,') ',False),expf(e,lambda q:(t(q,'−',False),t(q,'i',True),t(q,'E',True),t(q,'(',False),t(q,'p',True),t(q,')',False),t(q,'t',True))),t(e,' ',False),t(e,'d',True),t(e,'p',True))),t(o,',      ',False),t(o,'E',True),t(o,'(',False),t(o,'p',True),t(o,') = ',False),frac(o,lambda n:sup(n,var('p'),plain('2')),lambda d:t(d,'2',False)),t(o,'.',False)))
set_display(find('ϕ(p) = 1/√(2π) ∫ dx Ψ(0,x) exp(-ipx).'), lambda o:(t(o,'ϕ',True),t(o,'(',False),t(o,'p',True),t(o,') = ',False),frac(o,lambda n:t(n,'1',False),lambda d:rad(d,lambda e:(t(e,'2',False),t(e,'π',False)))),t(o,' ',False),nary(o,'∫',lambda e:(t(e,'Ψ',True),t(e,'(',False),t(e,'0',False),t(e,',',False),t(e,'x',True),t(e,') ',False),expf(e,lambda q:(t(q,'−',False),t(q,'i',True),t(q,'p',True),t(q,'x',True))),t(e,' ',False),t(e,'d',True),t(e,'x',True))),t(o,'.',False)))

set_display(find('φj(x) = √(2/L) sin[jπ(x + L/2)/L],     -L/2 ≤ x ≤ L/2.'), lambda o:(phij(o),t(o,'(',False),t(o,'x',True),t(o,') = ',False),rad(o,lambda e:frac(e,lambda n:t(n,'2',False),lambda d:t(d,'L',True))),t(o,' ',False),func(o,'sin',lambda e:delim(e,'[',']',lambda q:frac(q,lambda n:(t(n,'j',True),t(n,'π',False),delim(n,'(',')',lambda z:(t(z,'x',True),t(z,' + ',False),frac(z,lambda a:t(a,'L',True),lambda b:t(b,'2',False))))),lambda d:t(d,'L',True)))),t(o,',      ',False),t(o,'−',False),frac(o,lambda n:t(n,'L',True),lambda d:t(d,'2',False)),t(o,' ≤ ',False),t(o,'x',True),t(o,' ≤ ',False),frac(o,lambda n:t(n,'L',True),lambda d:t(d,'2',False)),t(o,'.',False)))
set_display(find('ψn(x) = ∑j=1J cj(n) φj(x).'), lambda o:(psin(o),t(o,'(',False),t(o,'x',True),t(o,') = ',False),nary(o,'∑',lambda e:(subsup(e,var('c'),var('j'),lambda s:t(s,'(n)',False)),t(e,' ',False),phij(e),t(e,'(',False),t(e,'x',True),t(e,')',False)),lo=lambda s:(t(s,'j',True),t(s,'=1',False)),hi=lambda s:t(s,'J',True)),t(o,'.',False)))
set_display(find('Hᵢⱼ = Tᵢⱼ + Vᵢⱼ.'), lambda o:(Hij(o),t(o,' = ',False),Tij(o),t(o,' + ',False),Vij(o),t(o,'.',False)))
set_display(find('Tᵢⱼ = (j²π²/2L²) δᵢⱼ.'), lambda o:(Tij(o),t(o,' = ',False),frac(o,lambda n:(sup(n,var('j'),plain('2')),sup(n,plain('π'),plain('2'))),lambda d:(t(d,'2',False),sup(d,var('L'),plain('2')))),t(o,' ',False),deltaij(o),t(o,'.',False)))
set_display(find('Vᵢⱼ = ∫₋L/2ᴸ/2 φᵢ(x)V(x)φⱼ(x) dx,'), lambda o:(Vij(o),t(o,' = ',False),nary(o,'∫',lambda e:(phii(e),t(e,'(',False),t(e,'x',True),t(e,')',False),t(e,'V',True),t(e,'(',False),t(e,'x',True),t(e,')',False),phij(e),t(e,'(',False),t(e,'x',True),t(e,') ',False),t(e,'d',True),t(e,'x',True)),lo=lambda s:(t(s,'−',False),frac(s,lambda n:t(n,'L',True),lambda d:t(d,'2',False))),hi=lambda s:frac(s,lambda n:t(n,'L',True),lambda d:t(d,'2',False))),t(o,',',False)))
set_display(find('H c⁽ⁿ⁾ = Eₙ c⁽ⁿ⁾,'), lambda o:(t(o,'H',True),t(o,' ',False),sup(o,var('c'),lambda s:t(s,'(n)',False)),t(o,' = ',False),En(o),t(o,' ',False),sup(o,var('c'),lambda s:t(s,'(n)',False)),t(o,',',False)))

set_display(find('Hθ = exp(−2iθ)T + V[x exp(iθ)] = −(1/2) exp(−2iθ) d²/dx² + V[x exp(iθ)].'), lambda o:(Htheta(o),t(o,' = ',False),expf(o,lambda e:(t(e,'−2',False),t(e,'i',True),t(e,'θ',True))),t(o,'T + V',False),delim(o,'[',']',lambda e:(t(e,'x',True),t(e,' ',False),expf(e,lambda q:(t(q,'i',True),t(q,'θ',True))))),t(o,' = −',False),frac(o,lambda n:t(n,'1',False),lambda d:t(d,'2',False)),t(o,' ',False),expf(o,lambda e:(t(e,'−2',False),t(e,'i',True),t(e,'θ',True))),t(o,' ',False),frac(o,lambda n:sup(n,var('d'),plain('2')),lambda d:(t(d,'d',True),sup(d,var('x'),plain('2')))),t(o,' + V',False),delim(o,'[',']',lambda e:(t(e,'x',True),t(e,' ',False),expf(e,lambda q:(t(q,'i',True),t(q,'θ',True))))),t(o,'.',False)))
set_display(find('Hθ c⁽ⁿ⁾ = Eₙ(θ)c⁽ⁿ⁾.'), lambda o:(Htheta(o),t(o,' ',False),sup(o,var('c'),lambda s:t(s,'(n)',False)),t(o,' = ',False),En(o),t(o,'(',False),t(o,'θ',True),t(o,')',False),sup(o,var('c'),lambda s:t(s,'(n)',False)),t(o,'.',False)))
set_display(find('Eₙ = Eres,n − iΓₙ/2,     Γₙ = −2 Im(Eₙ).'), lambda o:(En(o),t(o,' = ',False),Eresn(o),t(o,' − ',False),t(o,'i',True),frac(o,lambda n:Gaman(n),lambda d:t(d,'2',False)),t(o,',      ',False),Gaman(o),t(o,' = −2 ',False),tnor(o,'Im'),delim(o,'(',')',lambda e:En(e)),t(o,'.',False)))
set_display(find('|T(E)|² ≈ (Γₙ/2)² / [(E−Eres,n)² + (Γₙ/2)²].'), lambda o:(sup(o,lambda e:delim(e,'|','|',lambda q:(t(q,'T',True),t(q,'(',False),t(q,'E',True),t(q,')',False))),plain('2')),t(o,' ≈ ',False),frac(o,lambda n:sup(n,lambda q:delim(q,'(',')',lambda z:frac(z,lambda a:Gaman(a),lambda b:t(b,'2',False))),plain('2')),lambda d:delim(d,'[',']',lambda q:(sup(q,lambda z:delim(z,'(',')',lambda a:(t(a,'E',True),t(a,' − ',False),Eresn(a))),plain('2')),t(q,' + ',False),sup(q,lambda z:delim(z,'(',')',lambda a:frac(a,lambda b:Gaman(b),lambda c:t(c,'2',False))),plain('2'))))),t(o,'.',False)))

doc.save(OUT)
print(OUT)

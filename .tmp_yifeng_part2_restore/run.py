from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pathlib import Path

src = Path('Yifeng CHEN 999016959/Quantum_Scattering_Project_Report.docx')
doc = Document(src)
paras = doc.paragraphs
heading_idx = next(i for i,p in enumerate(paras) if p.text == 'Part 2 — Regularized Dirac delta function')
img_idx = next(i for i,p in enumerate(paras[heading_idx+1:], start=heading_idx+1) if p._p.xpath('.//w:drawing'))
anchor = paras[img_idx]
style = paras[heading_idx+1].style

# Keep only the three core steps before Figure 4.
for p in list(doc.paragraphs[heading_idx+1:img_idx]):
    p._element.getparent().remove(p._element)

def add_para_before(anchor, text=None, center=False, after=2.0):
    el = OxmlElement('w:p')
    anchor._p.addprevious(el)
    p = Paragraph(el, anchor._parent)
    p.style = style
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    if text is not None:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
    return p

M='m:'
def me(tag,val=None):
    e=OxmlElement(M+tag)
    if val is not None: e.set(qn('m:val'),val)
    return e

def mr(text, italic=True):
    r=me('r'); rp=me('rPr'); rp.append(me('sty','i' if italic else 'p')); r.append(rp)
    t=me('t'); t.text=text; r.append(t); return r

def txt(parent,text,italic=True): parent.append(mr(text,italic))
def sub(parent,base,subtxt,base_italic=True,sub_italic=True):
    ss=me('sSub'); e=me('e'); txt(e,base,base_italic); s=me('sub'); txt(s,subtxt,sub_italic); ss.extend([e,s]); parent.append(ss)
def frac(parent,num_builder,den_builder):
    f=me('f'); n=me('num'); d=me('den'); num_builder(n); den_builder(d); f.extend([n,d]); parent.append(f)
def integral(parent,lo,hi,expr_builder):
    n=me('nary'); pr=me('naryPr'); pr.append(me('chr','∫')); pr.append(me('limLoc','undOvr')); n.append(pr)
    s=me('sub'); txt(s,lo,False); u=me('sup'); txt(u,hi,False); e=me('e'); expr_builder(e); n.extend([s,u,e]); parent.append(n)
def limit(parent,sub_builder,expr_builder):
    limlow=me('limLow'); e=me('e'); txt(e,'lim',False); lim=me('lim'); sub_builder(lim); limlow.extend([e,lim]); parent.append(limlow)
    txt(parent,' ',False); expr_builder(parent)
def set_math_para(p,builder):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    op=me('oMathPara'); o=me('oMath'); builder(o); op.append(o); p._p.append(op)

# 1) Test-function definition
add_para_before(anchor, 'The Dirac delta function is defined through its action on a smooth test function f(p):', after=1.0)
p = add_para_before(anchor, center=True, after=2.0)
def eq1(o):
    def expr(e):
        for a,it in [('f',1),('(',0),('p',1),(')',0),('δ',1),('(',0),('p',1),(')',0),('d',1),('p',1)]: txt(e,a,bool(it))
    integral(o,'−∞','+∞',expr)
    for a,it in [(' = ',0),('f',1),('(',0),('0',0),(')',0)]: txt(o,a,bool(it))
set_math_para(p,eq1)

# 2) Finite-L regularization used in the project
add_para_before(anchor, 'For the numerical calculation, the regularized Dirac delta function used in this project is', after=1.0)
p = add_para_before(anchor, center=True, after=2.0)
def eq2(o):
    sub(o,'δ','L',True,True)
    for a,it in [('(',0),('p',1),(')',0),(' = ',0)]: txt(o,a,bool(it))
    frac(o, lambda n: txt(n,'1',False), lambda d: (txt(d,'2',False),txt(d,'π',False)))
    txt(o,' ',False)
    def expr(e):
        s=me('sSup'); base=me('e'); txt(base,'e',True); sup=me('sup')
        for a,it in [('i',1),('p',1),('x',1)]: txt(sup,a,bool(it))
        s.extend([base,sup]); e.append(s); txt(e,'d',True); txt(e,'x',True)
    integral(o,'−L','L',expr)
    txt(o,' = ',False)
    def nb(n):
        f=me('func'); f.append(me('funcPr')); fn=me('fName'); txt(fn,'sin',False); arg=me('e')
        for a,it in [('(',0),('L',1),('p',1),(')',0)]: txt(arg,a,bool(it))
        f.extend([fn,arg]); n.append(f)
    frac(o, nb, lambda d: (txt(d,'π',False),txt(d,'p',True)))
set_math_para(p,eq2)

# 3) L -> infinity test-function limit
add_para_before(anchor, 'As L → ∞, the regularized function approaches the Dirac delta distribution, so that', after=1.0)
p = add_para_before(anchor, center=True, after=3.0)
def eq3(o):
    def subb(s):
        for a,it in [('L',1),(' → ',0),('∞',0)]: txt(s,a,bool(it))
    def expr(parent):
        def integrand(e):
            for a,it in [('f',1),('(',0),('p',1),(')',0)]: txt(e,a,bool(it))
            sub(e,'δ','L',True,True)
            for a,it in [('(',0),('p',1),(')',0),('d',1),('p',1)]: txt(e,a,bool(it))
        integral(parent,'−∞','+∞',integrand)
    limit(o,subb,expr)
    for a,it in [(' = ',0),('f',1),('(',0),('0',0),(')',0)]: txt(o,a,bool(it))
set_math_para(p,eq3)

doc.save(src)
print('Yifeng Part 2 reduced to the three core steps.')

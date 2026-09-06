from docx import Document
from pathlib import Path

src = Path('Shufan ZHANG 999018435/quantum_scattering_report.docx')
doc = Document(src)

repl = {
    'For a free particle, V(x)=0. Using m=ħ=1, the time-independent Schrödinger equation is':
    'For the free-particle case, V(x)=0. With m=ħ=1, the stationary Schrödinger equation reduces to',

    'With E=p²/2, the equation becomes d²uₚ(x)/dx²+p²uₚ(x)=0, whose solutions are plane waves exp(±ipx). The momentum eigenstate with momentum p used in the Fourier expansion is':
    'Substituting E=p²/2 gives d²uₚ(x)/dx²+p²uₚ(x)=0, so the solutions are plane waves exp(±ipx). The normalized momentum eigenfunction used for the expansion is',

    'The initial Gaussian wavepacket can therefore be expanded in these free-particle momentum eigenstates. Its momentum-space amplitude ϕ(p) is obtained from the Fourier transform,':
    'This plane-wave basis is then used to decompose the initial Gaussian wavepacket. The corresponding momentum-space coefficient ϕ(p) is obtained from',

    'Each momentum component then evolves with the free-particle phase exp(-ip²t/2), and the wavepacket at time t is':
    'Under free evolution, the coefficient ϕ(p) remains unchanged while each momentum component acquires the phase exp(-ip²t/2). Hence the wavepacket at time t is',

    'The Dirac delta function is defined through its action on a smooth test function f(p):':
    'For a smooth test function f(p), the Dirac delta is characterized by',

    'For the numerical calculation, the regularized Dirac delta function used in this project is':
    'The finite-L regularization used in the numerical calculation is',

    'As L → ∞, the regularized function approaches the Dirac delta distribution, so that':
    'In the large-L limit, this regularized form reproduces the defining action of the Dirac delta:',

    'For the numerical test, the smooth test function f(p) is chosen as the momentum-space Gaussian ϕ(p) from Part 1, with α=1 and p₀=5. The numerical integration window is fixed at p∈[-10,10] for every L.':
    'To check this behavior numerically, f(p) is taken as the momentum-space Gaussian ϕ(p) from Part 1, with α=1 and p₀=5. The integration window is kept fixed at p∈[-10,10] for all L.'
}

found = set()
for p in doc.paragraphs:
    if p.text in repl:
        old = p.text
        fmt = None
        if p.runs:
            r0 = p.runs[0]
            fmt = (r0.font.name, r0.font.size, r0.bold, r0.italic)
        for r in list(p.runs):
            p._p.remove(r._r)
        r = p.add_run(repl[old])
        if fmt:
            name, size, bold, italic = fmt
            if name:
                r.font.name = name
            if size:
                r.font.size = size
            r.bold = bold
            r.italic = italic
        found.add(old)

missing = set(repl) - found
if missing:
    raise RuntimeError(f'Missing expected paragraphs: {missing}')

doc.save(src)
print(f'Replaced {len(found)} transition paragraphs.')
